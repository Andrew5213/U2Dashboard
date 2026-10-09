"""Leitura das planilhas de cronograma PMEx (um ficheiro por disciplina).

Junta os ficheiros de uma província numa única estrutura: os "marcos externos"
que ligavam uma disciplina à outra viram dependências comuns entre tarefas.
Usado pelo importador (scripts/import_schedule.py) e para gerar os dados de teste.
"""
from __future__ import annotations

import warnings
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl

from src.services.schedule_engine import CAL_CORRIDO, CAL_UTIL

_FIRST_HOLIDAY_ROW, _LAST_HOLIDAY_ROW = 28, 57
_FIRST_MILESTONE_ROW = 21
_FIRST_EXTERNAL_ROW, _LAST_EXTERNAL_ROW = 16, 25


@dataclass(frozen=True)
class XlsxTask:
    key: str                    # "<disciplina>:<ID da linha>", único entre os ficheiros
    name: str
    level: int                  # 1 = disciplina, 2 = grupo, 3 = tarefa
    parent_key: str | None
    is_summary: bool
    duration: float | None      # None = ainda não se sabe (fica em branco no ClickUp)
    calendar: str
    percent: float              # 0..100
    predecessors: tuple[str, ...]
    resources: tuple[str, ...]
    expected_start: date | None     # valores calculados pela planilha (para conferência)
    expected_finish: date | None
    expected_pos_start: float | None
    expected_pos_finish: float | None


@dataclass(frozen=True)
class XlsxSchedule:
    province: str
    project: str
    start: date
    weekend_mask: str
    holidays: tuple[tuple[date, str], ...]
    tasks: tuple[XlsxTask, ...] = field(default_factory=tuple)


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _split(value, separator: str) -> list[str]:
    return [part.strip() for part in str(value or "").split(separator) if part.strip()]


def load_schedule(paths: list[str]) -> XlsxSchedule:
    """Lê os ficheiros na ordem dada (a ordem define a ordem das disciplinas)."""
    books = []
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # extensões de formatação que o openpyxl não lê
        for path in paths:
            books.append(openpyxl.load_workbook(path, data_only=True))

    # Marcos fornecidos (aba Resumo): código → tarefa que o entrega
    milestones: dict[str, str] = {}
    disciplines: list[str] = []
    for book in books:
        discipline = str(book["Parâmetros"]["B4"].value).strip()
        disciplines.append(discipline)
        summary = book["Resumo"]
        for row in range(_FIRST_MILESTONE_ROW, summary.max_row + 1):
            code, task_id = summary.cell(row, 1).value, summary.cell(row, 3).value
            if code and isinstance(task_id, (int, float)):
                milestones[str(code).strip()] = f"{discipline}:{int(task_id)}"

    tasks: list[XlsxTask] = []
    holidays: dict[date, str] = {}
    for book, discipline in zip(books, disciplines):
        params = book["Parâmetros"]
        for row in range(_FIRST_HOLIDAY_ROW, _LAST_HOLIDAY_ROW + 1):
            day = _as_date(params.cell(row, 1).value)
            if day:
                holidays[day] = str(params.cell(row, 2).value or "").strip()

        # Marcos que esta disciplina recebe de fora (aba Parâmetros): código → descrição
        declared = {
            str(params.cell(row, 1).value).strip(): str(params.cell(row, 2).value or "").strip()
            for row in range(_FIRST_EXTERNAL_ROW, _LAST_EXTERNAL_ROW + 1) if params.cell(row, 1).value
        }

        sheet = book["Cronograma"]
        for row in range(2, sheet.max_row + 1):
            row_id = sheet.cell(row, 1).value
            if row_id is None:
                continue
            cell = lambda column: sheet[f"{column}{row}"].value  # noqa: E731
            is_summary = str(cell("C")).strip() == "Resumo"
            predecessors = [f"{discipline}:{int(p)}" for p in _split(cell("I"), ";")]
            parent = cell("E")
            parent_key = f"{discipline}:{int(parent)}" if parent else None
            for code in _split(cell("J"), ";"):
                if code not in milestones:
                    if code not in declared:
                        raise ValueError(f"{discipline} linha {row}: marco {code} não é fornecido por nenhum ficheiro")
                    # Marco de fora do projeto (ex.: "Torre entregue no site", do fornecedor):
                    # vira uma tarefa própria logo antes de quem depende dela, sem duração —
                    # o prazo de entrega não se sabe; preenche-se no ClickUp.
                    milestones[code] = f"{discipline}:{code}"
                    tasks.append(XlsxTask(
                        key=milestones[code], name=declared[code].split(" (")[0] or code,
                        level=int(cell("D")), parent_key=parent_key, is_summary=False,
                        duration=None, calendar=CAL_UTIL, percent=0.0, predecessors=(), resources=(),
                        expected_start=None, expected_finish=None,
                        expected_pos_start=None, expected_pos_finish=None,
                    ))
                predecessors.append(milestones[code])
            tasks.append(XlsxTask(
                key=f"{discipline}:{int(row_id)}",
                name=str(cell("B")).strip(),
                level=int(cell("D")),
                parent_key=parent_key,
                is_summary=is_summary,
                duration=0.0 if is_summary else float(cell("F") or 0),
                calendar=CAL_CORRIDO if str(cell("L") or "").strip() == "Corrido" else CAL_UTIL,
                percent=0.0 if is_summary else round(float(cell("P") or 0) * 100, 4),
                predecessors=tuple(dict.fromkeys(predecessors)),
                resources=tuple(_split(cell("K"), "&")),
                expected_start=_as_date(cell("G")),
                expected_finish=_as_date(cell("H")),
                expected_pos_start=None if is_summary else float(cell("N")),
                expected_pos_finish=None if is_summary else float(cell("O")),
            ))

    first = books[0]["Parâmetros"]
    return XlsxSchedule(
        province=str(first["B3"].value).strip(),
        project=str(first["B5"].value).strip(),
        start=_as_date(first["B6"].value),
        weekend_mask=str(first["B9"].value).strip(),
        holidays=tuple(sorted(holidays.items())),
        tasks=tuple(tasks),
    )
