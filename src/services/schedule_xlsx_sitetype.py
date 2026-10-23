"""Leitura das planilhas de tipos de site de transmissão ("Transmitter Site Type 1–4 - Work
Plan", 2026-10-21): uma aba "Site FM - Type N" com a WBS do site e a aba "Parameters".

A WBS vem em linhas com ID, tipo (Project/Stage/Task), nível e pai. No ClickUp:
    Project (nível 1)            → a própria lista (não vira tarefa)
    Stage de nível 2             → disciplina
    Stage de nível 3             → grupo
    Task de nível 2, 3 ou 4      → tarefa (de topo, na disciplina ou no grupo)
A coluna "Weight (%)" é o peso do item dentro do seu pai e vira o campo "Peso (%)".

As datas não vêm: o cronograma parte do início que alguém puser na primeira linha, como nas
outras listas. Mas a ordem das etapas, na planilha, só está nas datas de "Stage Start" (digitadas
à mão, sem antecessoras entre etapas) — então cada etapa passa a depender da anterior na planilha
quando esta começou antes (decisão de 2026-10-23); as que começam juntas (Obra Civil e Logística)
correm em paralelo. Para uma etapa começar num dia certo, preenche-se o "Início fixado" dela.
"""
from __future__ import annotations

import re
import warnings
from datetime import date, datetime, time, timedelta

import openpyxl

from src.services.schedule_engine import CAL_CORRIDO, CAL_UTIL

SHEET_PREFIX = "Site FM - Type"
_LEVEL_OF = {("Stage", 2): 1, ("Stage", 3): 2, ("Task", 2): 1, ("Task", 3): 2, ("Task", 4): 3}


def is_site_type_workbook(book) -> bool:
    return any(name.startswith(SHEET_PREFIX) for name in book.sheetnames)


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _ids(value) -> list[int]:
    return [int(part) for part in re.split(r"[;,\s]+", str(value if value is not None else "")) if part.isdigit()]


def _resources(value) -> tuple[str, ...]:
    return tuple(part.strip() for part in str(value or "").split("&") if part.strip())


def _clock(value) -> time | None:
    if isinstance(value, datetime):
        return value.time()
    return value if isinstance(value, time) else None


def _parameters(book) -> tuple[str, tuple[tuple[date, str], ...], tuple[str, ...]]:
    """Folga semanal, feriados e a jornada (linhas "jornada=" e "almoço=" da descrição da lista)."""
    mask, in_holidays = "0000001", False
    holidays: list[tuple[date, str]] = []
    start = lunch = None
    hours = lunch_hours = 0.0
    for row in book["Parameters"].iter_rows(values_only=True):
        label = str(row[0] or "").lower()
        if label.startswith("weekend pattern") and row[1]:
            mask = str(row[1]).strip()
        elif label.startswith("workday start"):
            start = _clock(row[1])
        elif label.startswith("working hours per day"):
            hours = float(row[1] or 0)
        elif label.startswith("lunch break start"):
            lunch = _clock(row[1])
        elif label.startswith("lunch break duration"):
            lunch_hours = float(row[1] or 0)
        elif label == "date":
            in_holidays = True
        elif in_holidays and (day := _as_date(row[0])):
            holidays.append((day, str(row[1] or "").strip()))
    lines: tuple[str, ...] = ()
    if start and hours:
        def at(clock: time, plus_hours: float) -> str:
            return (datetime.combine(date(2000, 1, 1), clock) + timedelta(hours=plus_hours)).strftime("%H:%M")
        lines = (f"jornada={start:%H:%M}-{at(start, hours + (lunch_hours if lunch else 0))}",)
        if lunch and lunch_hours:
            lines += (f"almoço={lunch:%H:%M}-{at(lunch, lunch_hours)}",)
    return mask, tuple(sorted(holidays)), lines


def load_site_type_schedule(path: str):
    from src.services.schedule_xlsx import XlsxSchedule, XlsxTask   # evita import circular

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")   # extensões de formatação que o openpyxl não lê
        book = openpyxl.load_workbook(path, data_only=True)
    sheet = next(book[name] for name in book.sheetnames if name.startswith(SHEET_PREFIX))
    header = [str(c.value or "").strip() for c in sheet[1]]
    col = {name: header.index(name) for name in header if name}

    rows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[col["ID"]] is not None]
    kinds = {int(r[col["ID"]]): (str(r[col["Type"]]), int(r[col["WBS Level"]])) for r in rows}
    key = lambda row_id: f"T:{row_id}"                                         # noqa: E731
    project = next((str(r[col["Name"]]) for r in rows if str(r[col["Type"]]) == "Project"), sheet.title)

    mask, holidays, lines = _parameters(book)
    stages = [r for r in rows if kinds[int(r[col["ID"]])] == ("Stage", 2)]
    stage_starts = {int(r[col["ID"]]): r[col["Stage Start"]] for r in stages
                    if isinstance(r[col["Stage Start"]], datetime)}
    # cada etapa espera a anterior na planilha, se esta começou antes (as que começam juntas: paralelo)
    chain: dict[int, int] = {}
    for before, after in zip(stages, stages[1:]):
        a, b = int(before[col["ID"]]), int(after[col["ID"]])
        if a in stage_starts and b in stage_starts and stage_starts[a] < stage_starts[b]:
            chain[b] = a

    tasks: list[XlsxTask] = []
    for r in rows:
        row_id, name = int(r[col["ID"]]), str(r[col["Name"]] or "").strip()
        kind, level = kinds[row_id]
        if kind == "Project":
            continue
        if (kind, level) not in _LEVEL_OF:
            raise ValueError(f"{sheet.title} ID {row_id} ({name}): {kind} no nível {level} — "
                             "o ClickUp só tem 3 níveis abaixo da lista")
        parent = int(r[col["Parent ID"]]) if r[col["Parent ID"]] is not None else None
        parent_key = key(parent) if parent is not None and kinds.get(parent, ("", 0))[0] != "Project" else None
        is_summary = kind == "Stage"
        weight = r[col["Weight (%)"]]
        tasks.append(XlsxTask(
            key=key(row_id), name=name, level=_LEVEL_OF[(kind, level)], parent_key=parent_key,
            is_summary=is_summary,
            duration=0.0 if is_summary else float(r[col["Task Duration (days)"]] or 0.0),
            calendar=CAL_CORRIDO if str(r[col["Calendar"]] or "").strip().lower() == "calendar" else CAL_UTIL,
            percent=0.0,
            predecessors=tuple(key(p) for p in _ids(r[col["Predecessors"]]) + ([chain[row_id]] if row_id in chain else [])
                               if p != row_id),
            resources=_resources(r[col["Resource Names"]]),
            expected_start=None, expected_finish=None, expected_pos_start=None, expected_pos_finish=None,
            weight=round(float(weight) * 100, 2) if weight is not None else None,
            source_id=row_id, wbs=str(_LEVEL_OF[(kind, level)]),
        ))

    return XlsxSchedule(province=sheet.title.replace(SHEET_PREFIX, "Type").strip(), project=project,
                        start=None, weekend_mask=mask, holidays=holidays, tasks=tuple(tasks), list_lines=lines)
