"""Leitura das planilhas PMEx dos estúdios (formato diferente das dos sites FM).

Cada província tem dois ficheiros — Obra Civil e Área Técnica — e cada um tem uma
aba por sala: STUDIO, CTP e NEWSROOM. Cada aba é uma tabela no estilo do MS
Project (ID, Nome, Tipo, Nível WBS, ID Pai, Duração, Início, Fim, Antecessores,
Recursos), com até quatro níveis: projeto → fase → grupo → tarefa. A semana é de
segunda a sexta (WORKDAY do Excel) e os feriados ficam em STUDIO!L4:L10.

Cada sala é uma **lista própria** no ClickUp (STUDIO, CTP, NEWSROOM) — nenhuma
dependência atravessa de uma sala para outra. Dentro da lista, o mesmo padrão dos
sites FM, em três níveis:
  1. **disciplina** = o ficheiro (Obra Civil, Área Técnica);
  2. **grupo** = o resumo mais fundo acima das tarefas (Demolições e Remoções,
     Studio Equipment Installation, Rack Installation…); as fases que só agrupam
     outros resumos ("Civil Works", "Technical Center") são achatadas;
  3. **tarefa**.
A linha de nível 1 (o projeto da aba) e os marcos de duração zero ("Fim da Obra
Civil", "End of Work") não viram tarefas. Os "Marco externo — Obra Civil concluída:
<tarefa>" da Área Técnica viram dependências comuns da tarefa da Obra Civil com
esse nome, na mesma sala.
"""
from __future__ import annotations

import re
import unicodedata
import warnings
from dataclasses import dataclass
from datetime import date, datetime

import openpyxl

from src.services.schedule_engine import CAL_UTIL

ROOMS = ("STUDIO", "CTP", "NEWSROOM")
WEEKEND_MASK = "0000011"            # WORKDAY do Excel sem 3º argumento: sábado e domingo
_HOLIDAY_ROWS = range(4, 11)        # STUDIO!L4:L10 (nome definido "Feriados")
_EXTERNAL_PREFIX = re.compile(r"^marco externo\b.*?:\s*", re.I)


@dataclass(frozen=True)
class _Row:
    sheet: str
    discipline: str
    id: int
    name: str
    kind: str               # "Resumo" | "Task" | "Marco"
    level: int
    parent: int | None
    duration: float | None
    predecessors: tuple[int, ...]
    resources: tuple[str, ...]


def is_studio_workbook(book) -> bool:
    return "Parâmetros" not in book.sheetnames and all(room in book.sheetnames for room in ROOMS)


def _norm(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", text or "")
    return re.sub(r"\s+", " ", nfkd.encode("ascii", "ignore").decode("ascii")).strip().lower()


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _discipline(path: str) -> str:
    name = path.replace("\\", "/").rsplit("/", 1)[-1].lower()
    if "obra_civil" in name:
        return "Obra Civil"
    if "area_tecnica" in name:
        return "Área Técnica"
    raise ValueError(f"não sei a disciplina do ficheiro {path} (esperado *_Obra_Civil / *_Area_Tecnica)")


def _ints(value) -> tuple[int, ...]:
    if value in (None, ""):
        return ()
    if isinstance(value, (int, float)):
        return (int(value),)
    return tuple(int(float(part)) for part in str(value).replace(",", ";").split(";") if part.strip())


def _read_rows(book, discipline: str, room: str) -> list[_Row]:
    sheet = book[room]
    rows: list[_Row] = []
    for r in range(2, sheet.max_row + 1):
        row_id = sheet.cell(r, 1).value
        if row_id in (None, ""):
            continue
        kind = str(sheet.cell(r, 3).value or "").strip()
        duration = sheet.cell(r, 6).value
        rows.append(_Row(
            sheet=room, discipline=discipline, id=int(row_id),
            name=re.sub(r"\s+", " ", str(sheet.cell(r, 2).value or "")).strip(),
            kind=kind, level=int(sheet.cell(r, 4).value or 0),
            parent=int(sheet.cell(r, 5).value) if sheet.cell(r, 5).value not in (None, "") else None,
            duration=float(duration) if isinstance(duration, (int, float)) else None,
            predecessors=_ints(sheet.cell(r, 9).value),
            resources=tuple(p.strip() for p in str(sheet.cell(r, 10).value or "").split("&") if p.strip()),
        ))
    return rows


def _check_formulas(book, discipline: str, room: str, rows: list[_Row]) -> None:
    """A coluna Antecessores tem de bater com as linhas que a fórmula de Início usa —
    é o que garante que as dependências importadas são as que a planilha calcula."""
    sheet = book[room]
    by_id = {row.id: row for row in rows}
    for r in range(2, sheet.max_row + 1):
        formula = sheet.cell(r, 7).value
        row_id = sheet.cell(r, 1).value
        if row_id in (None, "") or not isinstance(formula, str) or "WORKDAY(" not in formula:
            continue
        referenced = {int(n) - 1 for n in re.findall(r"\bH(\d+)\b", formula)}
        declared = set(by_id[int(row_id)].predecessors)
        if referenced and referenced != declared:
            raise ValueError(f"{discipline}/{room} ID {row_id}: fórmula usa {sorted(referenced)}, "
                             f"Antecessores diz {sorted(declared)}")


def load_studio_schedule(paths: list[str], room: str | None = None):
    """Lê os ficheiros de estúdio de uma província e devolve o XlsxSchedule de UMA sala
    (STUDIO, CTP ou NEWSROOM) — cada sala vai para uma lista própria."""
    from src.services.schedule_xlsx import XlsxSchedule, XlsxTask   # evita import circular

    room = (room or "").strip().upper()
    if room not in ROOMS:
        raise ValueError(f"planilha de estúdio: informe a sala ({', '.join(ROOMS)}) — cada uma vai para uma lista")

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        formula_books = [openpyxl.load_workbook(p) for p in paths]
    disciplines = [_discipline(p) for p in paths]
    order = sorted(range(len(paths)), key=lambda i: 0 if disciplines[i] == "Obra Civil" else 1)

    first = formula_books[order[0]]["STUDIO"]
    start = _as_date(first["M1"].value)
    holidays: dict[date, str] = {}
    root_name = str(first["B2"].value or "")
    province = root_name.split("/")[0].replace("PMEx - RNA", "").strip() or "?"

    rows_by: dict[tuple[str, str], list[_Row]] = {}
    for i in order:
        book, discipline = formula_books[i], disciplines[i]
        sheet = book["STUDIO"]
        for r in _HOLIDAY_ROWS:
            day = _as_date(sheet.cell(r, 12).value)
            if day:
                holidays[day] = str(sheet.cell(r, 13).value or "").strip()
        rows = _read_rows(book, discipline, room)
        _check_formulas(book, discipline, room, rows)
        rows_by[(discipline, room)] = rows

    tasks: list[XlsxTask] = []
    if True:                                   # uma sala por lista
        # chave de cada linha que vira tarefa, e como resolver as que não viram
        keep: dict[tuple[str, int], str] = {}
        rows_in_room = {(row.discipline, row.id): row for d in ("Obra Civil", "Área Técnica")
                        for row in rows_by.get((d, room), [])}
        children: dict[tuple[str, int], list[_Row]] = {}
        for row in rows_in_room.values():
            if row.parent is not None:
                children.setdefault((row.discipline, row.parent), []).append(row)

        def is_group(row: _Row) -> bool:
            return row.kind == "Resumo" and row.level > 1 and any(
                c.kind != "Resumo" for c in children.get((row.discipline, row.id), []))

        def nearest_group(row: _Row) -> str:
            """Pai no ClickUp: o grupo mais próximo acima, ou a disciplina (o ficheiro)."""
            parent = rows_in_room.get((row.discipline, row.parent)) if row.parent is not None else None
            while parent is not None and not is_group(parent):
                parent = rows_in_room.get((parent.discipline, parent.parent)) if parent.parent is not None else None
            return keep[(parent.discipline, parent.id)] if parent is not None else row.discipline

        for discipline in ("Obra Civil", "Área Técnica"):
            for row in rows_by.get((discipline, room), []):
                if row.level <= 1 or row.kind == "Marco":
                    continue
                if row.kind == "Resumo" and not is_group(row):
                    continue                   # fase que só agrupa outros resumos: achatada
                keep[(row.discipline, row.id)] = f"{room}:{discipline}:{row.id}"
                if is_group(row) and nearest_group(row) != row.discipline:
                    raise ValueError(f"{discipline}/{room}: grupo dentro de grupo ({row.name}) — o ClickUp só tem 3 níveis")

        oc_by_name = {_norm(row.name): row for row in rows_by.get(("Obra Civil", room), [])}

        def resolve(row: _Row, depth: int = 0) -> list[str]:
            """Chaves que satisfazem uma dependência para `row` (que pode não ter virado tarefa)."""
            if depth > 20:
                raise ValueError(f"{room}: ciclo ao resolver dependências ({row.name})")
            if (row.discipline, row.id) in keep:
                return [keep[(row.discipline, row.id)]]
            if row.kind == "Marco":
                external = _EXTERNAL_PREFIX.sub("", row.name) if _EXTERNAL_PREFIX.match(row.name) else None
                if external is not None and row.discipline == "Área Técnica":
                    target = oc_by_name.get(_norm(external))
                    if target is None:
                        raise ValueError(f"{room}: marco externo sem tarefa correspondente na Obra Civil: {row.name}")
                    return resolve(target, depth + 1)
                return [k for p in row.predecessors for k in resolve(rows_in_room[(row.discipline, p)], depth + 1)]
            # fase achatada: depender dela é depender de tudo o que ela agrupa
            return [k for c in children.get((row.discipline, row.id), []) for k in resolve(c, depth + 1)]

        for discipline in ("Obra Civil", "Área Técnica"):
            if (discipline, room) not in rows_by:
                continue
            tasks.append(XlsxTask(
                key=discipline, name=discipline, level=1, parent_key=None, is_summary=True, duration=0.0,
                calendar=CAL_UTIL, percent=0.0, predecessors=(), resources=(),
                expected_start=None, expected_finish=None, expected_pos_start=None, expected_pos_finish=None,
            ))
            for row in rows_by.get((discipline, room), []):
                key = keep.get((row.discipline, row.id))
                if key is None:
                    continue
                group = is_group(row)
                predecessors = [k for p in row.predecessors
                                for k in resolve(rows_in_room[(row.discipline, p)]) if k != key]
                tasks.append(XlsxTask(
                    key=key, name=row.name,
                    level=2 if group or nearest_group(row) == row.discipline else 3,
                    parent_key=nearest_group(row),
                    is_summary=group,
                    duration=0.0 if group else row.duration,
                    calendar=CAL_UTIL, percent=0.0,
                    predecessors=tuple(dict.fromkeys(predecessors)),
                    resources=row.resources,
                    expected_start=None, expected_finish=None,
                    expected_pos_start=None, expected_pos_finish=None,
                ))

    return XlsxSchedule(
        province=province,
        project=f"PMEx - RNA {province} — {room}",
        start=start,
        weekend_mask=WEEKEND_MASK,
        holidays=tuple(sorted(holidays.items())),
        tasks=tuple(tasks),
    )
