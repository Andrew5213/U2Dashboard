"""Relatórios (PDF/XLSX) e listagens sobre uma lista de cronograma de 3 níveis."""
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from src.core.database import Base
from src.repositories.cache_repository import (
    CacheRepository, _FIELD_DATA_CONCLUSAO_ID, _FIELD_VENCIMENTO_ID,
)
from src.services.report_service import (
    DisciplineReportService, ListReportService, PeriodicReportService, ProvinceReportService,
    _fmt_progress, _status_with_progress,
)
from src.services.report_xlsx_service import render_periodic_xlsx
from src.services.schedule_fields import date_to_ms

TODAY = datetime.now(timezone.utc).date()
CALENDAR_OPTIONS = [{"id": "u", "name": "Útil", "orderindex": 0}, {"id": "c", "name": "Corrido", "orderindex": 1}]


def _ms(moment: datetime) -> str:
    return str(int(moment.replace(tzinfo=timezone.utc).timestamp() * 1000))


def schedule_task(task_id, parent=None, percent=0, duration=None, corrido=False, status="planejando",
                  status_type="open", start=None, due=None, closed=None, updated=None, created=0, team=None):
    return {
        "id": task_id, "name": task_id, "parent": parent, "list": {"id": "L1"},
        "status": {"status": status, "type": status_type, "color": "#ff7800"},
        "start_date": str(date_to_ms(start)) if start else None,
        "due_date": str(date_to_ms(due)) if due else None,
        "date_created": str(1_780_000_000_000 + created),
        "date_updated": _ms(updated) if updated else None,
        "group_assignees": [{"id": "g", "name": team}] if team else [],
        "custom_fields": [
            {"id": "p", "name": "% Concluído", "type": "manual_progress",
             "value": {"current": str(percent), "percent_completed": percent / 100}},
            {"id": "d", "name": "Duração (dias)", "type": "number",
             "value": None if duration is None else str(duration)},
            {"id": "k", "name": "Calendário", "type": "drop_down",
             "type_config": {"options": CALENDAR_OPTIONS}, "value": 1 if corrido else 0},
            {"id": _FIELD_DATA_CONCLUSAO_ID, "name": "Data de Conclusão", "type": "date",
             "value": _ms(closed) if closed else None},
        ],
    }


def plain_task(task_id, parent=None, vencimento=None, created=0):
    return {
        "id": task_id, "name": task_id, "parent": parent, "list": {"id": "L2"},
        "status": {"status": "planejando", "type": "open"},
        "date_created": str(1_780_000_000_000 + created),
        "custom_fields": [
            {"id": _FIELD_VENCIMENTO_ID, "name": "Vencimento", "type": "date",
             "value": str(date_to_ms(vencimento)) if vencimento else None},
        ],
    }


@pytest.fixture
async def db() -> AsyncSession:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        yield session
    await engine.dispose()


@pytest.fixture
async def repo(db) -> CacheRepository:
    """Site FM em miniatura (disciplina → grupo → tarefa) + uma lista comum de 2 níveis."""
    repo = CacheRepository(db)
    await repo.upsert_space({"id": "S", "name": "Space"})
    await repo.upsert_folder({"id": "F", "name": "LUBANGO"}, "S")
    await repo.upsert_list({"id": "L1", "name": "Site FM"}, "S", "F")
    await repo.upsert_list({"id": "L2", "name": "Estudios"}, "S", "F")
    now = datetime.utcnow()
    soon = TODAY + timedelta(days=3)
    for order, task in enumerate([
        schedule_task("Obra Civil", percent=50, duration=4, status="fazendo", status_type="custom",
                      start=TODAY - timedelta(days=2), due=soon, updated=now, team="PROEF"),
        schedule_task("Base do Gerador", "Obra Civil", percent=50, duration=4, status="fazendo",
                      status_type="custom", start=TODAY - timedelta(days=2), due=soon, updated=now),
        schedule_task("Concretagem", "Base do Gerador", percent=100, duration=2, status="complete",
                      status_type="done", start=TODAY - timedelta(days=2), due=TODAY - timedelta(days=1),
                      closed=now - timedelta(hours=2), updated=now, team="PROEF"),
        schedule_task("Cura do Concreto", "Base do Gerador", percent=40, duration=14, corrido=True,
                      status="fazendo", status_type="custom", start=TODAY, due=soon, updated=now, team="PROEF"),
        schedule_task("Logística", percent=0, duration=1.5, start=soon, due=soon),
    ]):
        await repo.upsert_task({**task, "date_created": str(1_780_000_000_000 + order)}, "L1")
    await repo.upsert_task(plain_task("Equipamentos", vencimento=soon), "L2")
    await repo.upsert_task(plain_task("Montagem", "Equipamentos", vencimento=soon, created=1), "L2")
    return repo


class TestListings:
    async def test_upcoming_lists_schedule_leaves_with_their_group(self, repo):
        names = [t["name"] for t in await repo.get_upcoming_tasks_by_folder("F", days=30)]
        assert "Base do Gerador / Cura do Concreto" in names
        assert "Logística" in names                 # disciplina sem filhos é a própria tarefa
        assert "Obra Civil" not in names            # resumo calculado
        assert "Base do Gerador" not in names

    async def test_plain_list_still_lists_the_top_level_task(self, repo):
        names = [t["name"] for t in await repo.get_upcoming_tasks_by_folder("F", days=30)]
        assert "Equipamentos" in names
        assert not any("Montagem" in name for name in names)

    async def test_space_listing_follows_the_same_rule(self, repo):
        names = [t["name"] for t in await repo.get_upcoming_tasks("S", 30)]
        assert "Base do Gerador / Cura do Concreto" in names and "Obra Civil" not in names


class TestPeriodUpdates:
    async def _tasks(self, repo) -> list[dict]:
        now = datetime.utcnow()
        folders = await repo.get_period_updates("S", now - timedelta(days=1), now + timedelta(minutes=1))
        return folders[0]["tasks"]

    async def test_leaves_nest_under_the_discipline_with_the_group_in_the_name(self, repo):
        tasks = await self._tasks(repo)
        assert {t["name"] for t in tasks} == {"Obra Civil"}
        by_category = {t["category"]: [s["name"] for s in t["subtasks"]] for t in tasks}
        assert by_category == {
            "concluded": ["Base do Gerador / Concretagem"],
            "updated": ["Base do Gerador / Cura do Concreto"],
        }

    async def test_summary_rewritten_by_the_engine_is_not_an_update_on_its_own(self, repo):
        # as tarefas voltam a "planejando": no período só a disciplina e o grupo,
        # ambos em "fazendo", têm data de atualização
        for leaf in ("Concretagem", "Cura do Concreto"):
            await repo.upsert_task(schedule_task(leaf, "Base do Gerador", duration=2), "L1")
        now = datetime.utcnow()
        folders = await repo.get_period_updates("S", now - timedelta(days=1), now + timedelta(minutes=1))
        assert folders == []

    async def test_open_schedule_task_shows_its_percent_beside_the_status(self, repo):
        updated = next(t for t in await self._tasks(repo) if t["category"] == "updated")
        assert _status_with_progress(updated["subtasks"][0]) == "fazendo 40%"
        assert _status_with_progress({"status": "complete", "progress_pct": 100}) == "complete"
        assert _status_with_progress({"status": "fazendo", "progress_pct": None}) == "fazendo"

    async def test_periodic_pdf_and_xlsx_render(self, repo, db):
        service = PeriodicReportService(db)
        assert (await service.generate_daily_pdf("S")).startswith(b"%PDF")
        now = datetime.utcnow()
        data = await service._build_data("S", now - timedelta(days=1), now, "Diario", "hoje")
        cells = [c for row in openpyxl.load_workbook(BytesIO(render_periodic_xlsx(data))).active.iter_rows(values_only=True) for c in row]
        assert "fazendo 40%" in cells


class TestListReport:
    async def test_rows_carry_level_dates_duration_and_percent(self, repo, db):
        data = await ListReportService(db)._build_data("L1")
        assert data["is_schedule"] is True
        rows = {r["name"]: r for r in data["rows"]}
        assert [r["level"] for r in data["rows"]] == [0, 1, 2, 2, 0]
        assert rows["Obra Civil"]["is_summary"] and rows["Base do Gerador"]["is_summary"]
        cure = rows["Cura do Concreto"]
        assert not cure["is_summary"]
        assert cure["duration_fmt"] == "14 d corr."
        assert cure["start_fmt"] == TODAY.strftime("%d/%m/%Y")
        assert cure["progress_pct"] == 40
        assert cure["assignees_str"] == "PROEF"
        assert rows["Logística"]["duration_fmt"] == "1,5 d"

    async def test_only_leaves_count_as_tasks_and_the_header_sums_the_list_up(self, repo, db):
        data = await ListReportService(db)._build_data("L1")
        assert data["total_disciplines"] == 2
        assert data["total_activities"] == 3        # 2 tarefas + a disciplina sem filhos
        start = (TODAY - timedelta(days=2)).strftime("%d/%m/%Y")
        end = (TODAY + timedelta(days=3)).strftime("%d/%m/%Y")
        assert start in data["schedule_summary"] and end in data["schedule_summary"]

    async def test_english_report_uses_dot_and_calendar_days(self, repo, db):
        rows = {r["name"]: r for r in (await ListReportService(db)._build_data("L1", lang="en"))["rows"]}
        # o nome também sai em inglês: é o mesmo dicionário da tela
        assert rows["Concrete Curing"]["duration_fmt"] == "14 cal. d"
        assert rows["Logística"]["duration_fmt"] == "1.5 d"

    async def test_list_without_dates_renders_blank_columns(self, db):
        repo = CacheRepository(db)
        await repo.upsert_space({"id": "S", "name": "Space"})
        await repo.upsert_folder({"id": "F", "name": "LUENA"}, "S")
        await repo.upsert_list({"id": "L1", "name": "Site FM"}, "S", "F")
        await repo.upsert_task(schedule_task("Obra Civil"), "L1")
        await repo.upsert_task(schedule_task("Fundação", "Obra Civil", duration=2, created=1), "L1")
        data = await ListReportService(db)._build_data("L1")
        assert data["rows"][1]["start_fmt"] == "-" and data["rows"][1]["end_fmt"] == "-"
        assert (await ListReportService(db).generate_pdf("L1")).startswith(b"%PDF")

    async def test_pdfs_render_for_list_and_discipline(self, repo, db):
        assert (await ListReportService(db).generate_pdf("L1")).startswith(b"%PDF")
        assert (await DisciplineReportService(db).generate_pdf("Obra Civil")).startswith(b"%PDF")
        assert (await DisciplineReportService(db).generate_pdf("Base do Gerador")).startswith(b"%PDF")

    async def test_discipline_report_counts_leaves_below_the_groups(self, repo, db):
        data = await DisciplineReportService(db)._build_data("Obra Civil")
        assert data["is_schedule"] is True and data["total_activities"] == 2
        assert [r["level"] for r in data["rows"]] == [0, 1, 2, 2]

    async def test_plain_list_keeps_the_old_layout(self, repo, db):
        data = await ListReportService(db)._build_data("L2")
        assert data["is_schedule"] is False and data["schedule_summary"] is None
        assert "level" not in data["rows"][0]
        assert (await ListReportService(db).generate_pdf("L2")).startswith(b"%PDF")


class TestProvinceReport:
    async def test_schedule_list_gets_schedule_columns_and_the_plain_list_does_not(self, repo, db):
        data = await ProvinceReportService(db)._build_data("F")
        lists = {lst["name"]: lst for lst in data["lists_detail"]}
        assert lists["Site FM"]["is_schedule"] and not lists["Estudios"]["is_schedule"]
        assert [(r["name"], r["level"]) for r in lists["Site FM"]["tasks"]] == [
            ("Obra Civil", 0), ("Base do Gerador", 1), ("Concretagem", 2), ("Cura do Concreto", 2), ("Logística", 0),
        ]
        assert "duration_fmt" not in lists["Estudios"]["tasks"][0]

    async def test_pdf_and_xlsx_render(self, repo, db):
        service = ProvinceReportService(db)
        assert (await service.generate_pdf("F")).startswith(b"%PDF")
        sheet = openpyxl.load_workbook(BytesIO(await service.generate_xlsx("F"))).active
        rows = [row for row in sheet.iter_rows(values_only=True) if row[1] in ("Disciplina", "Grupo", "Tarefa")]
        assert [(r[0].strip(), r[1]) for r in rows[:4]] == [
            ("Obra Civil", "Disciplina"), ("Base do Gerador", "Grupo"),
            ("Concretagem", "Tarefa"), ("Cura do Concreto", "Tarefa"),
        ]
        assert rows[3][2] == "14 d corr." and rows[3][5] == pytest.approx(0.4)


def test_progress_never_rounds_to_the_extremes():
    assert _fmt_progress(99.99) == "99%"
    assert _fmt_progress(0.2) == "1%"
    assert _fmt_progress(100) == "100%"
    assert _fmt_progress(0) == "0%"


class TestGanttReport:
    async def _data(self, db, folder_id="F", lang="pt") -> dict:
        from src.services.gantt_report_service import GanttReportService
        return await GanttReportService(db)._build_data(folder_id, lang)

    async def test_schedule_list_brings_every_level_with_its_dates(self, repo, db):
        data = await self._data(db)
        site = next(lst for lst in data["lists"] if lst["name"] == "Site FM")
        rows = {r["name"]: r for r in site["rows"]}
        assert [(r["name"], r["level"]) for r in site["rows"]] == [
            ("Obra Civil", 0), ("Base do Gerador", 1), ("Concretagem", 2), ("Cura do Concreto", 2), ("Logística", 0),
        ]
        assert rows["Obra Civil"]["is_summary"] and not rows["Cura do Concreto"]["is_summary"]
        assert (rows["Cura do Concreto"]["start"], rows["Cura do Concreto"]["end"]) == (TODAY, TODAY + timedelta(days=3))
        assert rows["Cura do Concreto"]["progress_pct"] == 40
        assert rows["Concretagem"]["is_done"] and not rows["Concretagem"]["is_overdue"]

    async def test_timeline_covers_whole_weeks_around_the_tasks(self, repo, db):
        data = await self._data(db)
        assert data["first_day"] == TODAY - timedelta(days=2)
        assert data["last_day"] == TODAY + timedelta(days=3)
        assert data["range_start"].weekday() == 0 and data["range_end"].weekday() == 6
        assert data["range_start"] <= data["first_day"] and data["range_end"] >= data["last_day"]
        assert data["total_tasks"] == 4             # 3 folhas do cronograma + Equipamentos

    async def test_plain_parent_without_dates_spans_its_children(self, db):
        repo = CacheRepository(db)
        await repo.upsert_space({"id": "S", "name": "Space"})
        await repo.upsert_folder({"id": "F", "name": "LUANDA"}, "S")
        await repo.upsert_list({"id": "L2", "name": "Estudio"}, "S", "F")
        await repo.upsert_list({"id": "L3", "name": "Vazia"}, "S", "F")
        await repo.upsert_task(plain_task("Obra Civil"), "L2")
        await repo.upsert_task(plain_task("Pintura", "Obra Civil", vencimento=TODAY - timedelta(days=5), created=1), "L2")
        await repo.upsert_task(plain_task("Piso", "Obra Civil", vencimento=TODAY + timedelta(days=2), created=2), "L2")
        await repo.upsert_task(plain_task("Sem data", "Obra Civil", created=3), "L2")
        await repo.upsert_task({**plain_task("Solta"), "list": {"id": "L3"}}, "L3")
        data = await self._data(db)
        rows = {r["name"]: r for r in data["lists"][0]["rows"]}
        assert set(rows) == {"Obra Civil", "Pintura", "Piso"}
        assert (rows["Obra Civil"]["start"], rows["Obra Civil"]["end"]) == (TODAY - timedelta(days=5), TODAY + timedelta(days=2))
        assert rows["Pintura"]["is_overdue"] and not rows["Piso"]["is_overdue"]
        assert data["undated_lists"] == ["Vazia"]

    async def test_pdf_renders_with_and_without_dates(self, repo, db):
        from src.services.gantt_report_service import GanttReportService
        assert (await GanttReportService(db).generate_pdf("F")).startswith(b"%PDF")
        assert (await GanttReportService(db).generate_pdf("F", lang="en")).startswith(b"%PDF")
        await repo.upsert_folder({"id": "F2", "name": "BENGUELA"}, "S")
        await repo.upsert_list({"id": "L9", "name": "Site FM"}, "S", "F2")
        await repo.upsert_task({**schedule_task("Obra Civil sem data"), "list": {"id": "L9"}}, "L9")
        data = await self._data(db, "F2")
        assert data["lists"] == [] and data["undated_lists"] == ["Site FM"]
        assert (await GanttReportService(db).generate_pdf("F2")).startswith(b"%PDF")

    async def test_long_schedule_renders_on_several_pages(self, db):
        from src.services.gantt_report_service import GanttReportService
        repo = CacheRepository(db)
        await repo.upsert_space({"id": "S", "name": "Space"})
        await repo.upsert_folder({"id": "F", "name": "LUENA"}, "S")
        await repo.upsert_list({"id": "L1", "name": "Site FM"}, "S", "F")
        await repo.upsert_task(schedule_task("Obra", start=TODAY, due=TODAY + timedelta(days=400)), "L1")
        for i in range(90):
            await repo.upsert_task(schedule_task(f"T{i}", "Obra", duration=3, start=TODAY + timedelta(days=4 * i),
                                                 due=TODAY + timedelta(days=4 * i + 3), created=i + 1), "L1")
        from src.services.gantt_report_service import _GanttReport
        data = await GanttReportService(db)._build_data("F")
        pdf = _GanttReport(data)
        pdf.add_page()
        pdf.build_title(data)
        pdf.build_lists(data["lists"])
        assert pdf.page_no() > 1                    # 91 linhas não cabem numa página
        assert pdf._label_every() > 1               # mais de um ano: o rótulo salta semanas, a grade não
        assert (await GanttReportService(db).generate_pdf("F")).startswith(b"%PDF")

    async def test_continuation_page_starts_below_the_letterhead(self, db):
        """O timbrado põe as logos no topo de toda página; a tabela da segunda em
        diante começava no topo e passava por cima delas."""
        from src.services.gantt_report_service import _BOTTOM, _TOP, GanttReportService, _GanttReport
        repo = CacheRepository(db)
        await repo.upsert_space({"id": "S", "name": "Space"})
        await repo.upsert_folder({"id": "F", "name": "LUENA"}, "S")
        await repo.upsert_list({"id": "L1", "name": "Site FM"}, "S", "F")
        await repo.upsert_task(schedule_task("Obra", start=TODAY, due=TODAY + timedelta(days=10)), "L1")
        await repo.upsert_task(schedule_task("T1", "Obra", duration=3, start=TODAY,
                                             due=TODAY + timedelta(days=3), created=1), "L1")

        pdf = _GanttReport(await GanttReportService(db)._build_data("F"))
        pdf.add_page()
        pdf.set_y(_BOTTOM)          # não cabe mais nada: a próxima linha quebra a página
        pdf._ensure_room(20)

        assert pdf.page_no() == 2
        assert pdf.get_y() >= _TOP

    async def test_unknown_folder_is_an_error(self, db):
        from src.services.gantt_report_service import GanttReportService
        with pytest.raises(ValueError):
            await GanttReportService(db).generate_pdf("nao-existe")


class TestScheduleDeviation:
    """Atraso em dias: término atual das tarefas em cache × linha de base."""

    async def _baseline(self, db, finish_by_task: dict) -> None:
        from src.models.schedule_models import ScheduleBaseline
        for task_id, finish in finish_by_task.items():
            db.add(ScheduleBaseline(task_id=task_id, list_id="L1", start=finish, finish=finish))
        await db.commit()

    async def test_list_and_folder_carry_the_delay_in_days(self, repo, db):
        await self._baseline(db, {"Cura do Concreto": TODAY + timedelta(days=1), "Concretagem": TODAY - timedelta(days=1)})
        lists = {lst["name"]: lst for lst in await repo.get_lists_with_metrics("F")}
        assert lists["Site FM"]["delay_days"] == 2            # termina em +3, plano em +1
        assert lists["Site FM"]["baseline_finish"] == TODAY + timedelta(days=1)
        assert lists["Site FM"]["schedule_finish"] == TODAY + timedelta(days=3)
        assert "delay_days" not in lists["Estudios"]
        folder = (await repo.get_folders_with_metrics("S"))[0]
        assert folder["delay_days"] == 2
        assert (await repo.get_list_kpis("L1"))["delay_days"] == 2

    async def test_without_baseline_nothing_is_reported(self, repo):
        assert "delay_days" not in (await repo.get_lists_with_metrics("F"))[0]
        assert "delay_days" not in (await repo.get_folders_with_metrics("S"))[0]

    async def test_each_task_shows_its_own_deviation(self, repo, db):
        from src.services.dashboard_service import DashboardService
        await self._baseline(db, {"Obra Civil": TODAY, "Cura do Concreto": TODAY + timedelta(days=5)})
        tasks = {t.name: t for t in await DashboardService(db).get_list_tasks("L1")}
        assert tasks["Obra Civil"].delay_days == 3 and tasks["Obra Civil"].baseline_due == TODAY
        assert tasks["Logística"].delay_days is None            # criada depois da linha de base
        detail = await DashboardService(db).get_task_detail("Base do Gerador")
        cure = next(s for s in detail.subtasks if s.name == "Cura do Concreto")
        assert cure.delay_days == -2                            # adiantada


class TestSearch:
    async def test_matches_every_word_ignoring_case_and_accents(self, repo):
        found = await repo.search_tasks("S", "CURA concreto")
        assert [t["name"] for t in found] == ["Cura do Concreto"]
        assert found[0]["parent_name"] == "Base do Gerador" and found[0]["folder_name"] == "LUBANGO"
        assert [t["name"] for t in await repo.search_tasks("S", "logistica")] == ["Logística"]

    async def test_reaches_every_list_and_level(self, repo):
        names = {t["name"] for t in await repo.search_tasks("S", "o")}
        assert {"Obra Civil", "Montagem", "Concretagem"} <= names

    async def test_blank_or_unmatched_query_finds_nothing(self, repo):
        assert await repo.search_tasks("S", "   ") == []
        assert await repo.search_tasks("S", "transmissor inexistente") == []


class TestTeamView:
    async def test_stats_count_what_is_running_and_what_comes_this_week(self, repo):
        stats = {s["assignee"]: s for s in await repo.get_assignee_task_stats("S")}
        proef = stats["PROEF"]
        assert (proef["open"], proef["completed"]) == (1, 1)
        assert proef["in_progress"] == 1 and proef["next_7_days"] == 1

    async def test_tasks_carry_start_group_and_started_flag(self, repo):
        tasks = {t["name"]: t for t in await repo.get_tasks_by_assignee("S", "PROEF")}
        cure = tasks["Cura do Concreto"]
        assert cure["is_started"] and cure["parent_name"] == "Base do Gerador"
        assert cure["start_date"].date() == TODAY
        assert not tasks["Concretagem"]["is_started"]           # concluída não está "em andamento"


class TestGanttSpreadsheet:
    """Gantt em .xlsx: as barras são células pintadas, uma coluna por dia."""

    async def _sheet(self, db, folder_id="F", lang="pt"):
        from src.services.gantt_report_service import GanttReportService
        return openpyxl.load_workbook(BytesIO(await GanttReportService(db).generate_xlsx(folder_id, lang=lang))).active

    @staticmethod
    def _color(cell) -> str:
        return (cell.fill.fgColor.rgb or "")[-6:] if cell.fill and cell.fill.fill_type else ""

    def _row_of(self, sheet, name):
        return next(r for r in range(1, sheet.max_row + 1) if sheet.cell(row=r, column=1).value == name)

    async def test_task_bar_covers_its_days_with_the_done_part_first(self, repo, db):
        sheet = await self._sheet(db)
        row = self._row_of(sheet, "Cura do Concreto")              # hoje → hoje+3, 40%
        assert sheet.cell(row=row, column=2).value.date() == TODAY
        assert sheet.cell(row=row, column=5).value == pytest.approx(0.4)
        bar = [self._color(sheet.cell(row=row, column=c)) for c in range(7, sheet.max_column + 1)]
        painted = [color for color in bar if color in ("3B82F6", "BFDBFE")]
        assert len(painted) == 4                                   # 4 dias, início e fim inclusive
        assert painted == ["3B82F6", "3B82F6", "BFDBFE", "BFDBFE"]  # 40% de 4 dias → 2 células

    async def test_summaries_and_completed_tasks_have_their_own_colors(self, repo, db):
        sheet = await self._sheet(db)
        colors = lambda name: {self._color(sheet.cell(row=self._row_of(sheet, name), column=c))
                               for c in range(7, sheet.max_column + 1)}
        assert "475569" in colors("Obra Civil")                    # resumo
        assert "16A34A" in colors("Concretagem")                   # concluída
        assert sheet.row_dimensions[self._row_of(sheet, "Cura do Concreto")].outlineLevel == 2

    async def test_header_is_frozen_and_weeks_are_labelled_by_their_monday(self, repo, db):
        sheet = await self._sheet(db)
        assert sheet.freeze_panes == "G7"
        monday = TODAY - timedelta(days=TODAY.weekday())
        weeks = [sheet.cell(row=5, column=c).value for c in range(7, sheet.max_column + 1)]
        assert monday.strftime("%d/%m") in weeks
        assert sheet.cell(row=4, column=1).value == "Tarefa"

    async def test_english_and_empty_province(self, repo, db):
        sheet = await self._sheet(db, lang="en")
        assert sheet.cell(row=4, column=1).value == "Task" and sheet.cell(row=4, column=2).value == "Start"
        await repo.upsert_folder({"id": "F2", "name": "BENGUELA"}, "S")
        await repo.upsert_list({"id": "L9", "name": "Site FM"}, "S", "F2")
        await repo.upsert_task({**schedule_task("Obra sem data"), "list": {"id": "L9"}}, "L9")
        empty = await self._sheet(db, "F2")
        assert "Nenhuma tarefa" in empty.cell(row=3, column=1).value

    async def test_long_schedule_switches_to_one_column_per_week(self, db):
        repo = CacheRepository(db)
        await repo.upsert_space({"id": "S", "name": "Space"})
        await repo.upsert_folder({"id": "F", "name": "LUENA"}, "S")
        await repo.upsert_list({"id": "L1", "name": "Site FM"}, "S", "F")
        await repo.upsert_task(schedule_task("Longa", duration=300, start=TODAY, due=TODAY + timedelta(days=400)), "L1")
        sheet = await self._sheet(db)
        assert sheet.max_column - 6 < 70            # ~58 semanas, não 400 dias
        assert sheet.freeze_panes == "G6"           # sem a linha dos dias


class TestGanttOverviewStart:
    async def test_province_starts_on_its_first_real_start_date(self, repo):
        row = next(r for r in await repo.get_gantt_overview("S") if r["folder_id"] == "F")
        assert row["start_date"][:10] == (TODAY - timedelta(days=2)).isoformat()   # não a data de criação


class TestGanttPauses:
    """Pausa no Gantt: a tarefa e as filhas levam o trecho zebrado (PDF, XLSX e dashboard)."""

    async def _paused(self, repo):
        pause_fields = [
            {"id": "pa", "name": "Pausa de", "type": "date", "value": str(date_to_ms(TODAY + timedelta(days=1)))},
            {"id": "pb", "name": "Pausa até", "type": "date", "value": str(date_to_ms(TODAY + timedelta(days=2)))},
            {"id": "pc", "name": "Motivo da pausa", "type": "short_text", "value": "chuva"},
            {"id": "pd", "name": "Histórico de pausas", "type": "text",
             "value": f"{TODAY - timedelta(days=2):%d/%m/%Y} – {TODAY - timedelta(days=2):%d/%m/%Y} · licença"},
        ]
        group = schedule_task("Base do Gerador", "Obra Civil", percent=50, duration=4, status="fazendo",
                              status_type="custom", start=TODAY - timedelta(days=2), due=TODAY + timedelta(days=3))
        group["custom_fields"] += pause_fields
        await repo.upsert_task({**group, "date_created": str(1_780_000_000_001)}, "L1")

    async def test_rows_carry_their_own_and_inherited_pauses(self, repo, db):
        from src.services.gantt_report_service import GanttReportService
        await self._paused(repo)
        data = await GanttReportService(db)._build_data("F")
        site = next(lst for lst in data["lists"] if lst["name"] == "Site FM")
        rows = {r["name"]: r for r in site["rows"]}
        spans = lambda name: [(p["start"], p["end"], p["reason"]) for p in rows[name]["pauses"]]
        expected = [(TODAY - timedelta(days=2), TODAY - timedelta(days=2), "licença"),
                    (TODAY + timedelta(days=1), TODAY + timedelta(days=2), "chuva")]
        assert spans("Base do Gerador") == expected
        assert spans("Cura do Concreto") == expected                  # herdada da tarefa-mãe
        assert all(p["own"] for p in rows["Base do Gerador"]["pauses"])
        assert not any(p["own"] for p in rows["Cura do Concreto"]["pauses"])
        assert spans("Obra Civil") == [] and spans("Logística") == []

    async def test_pdf_and_spreadsheet_draw_the_pause(self, repo, db):
        from src.services.gantt_report_service import GanttReportService
        await self._paused(repo)
        assert (await GanttReportService(db).generate_pdf("F")).startswith(b"%PDF")
        sheet = openpyxl.load_workbook(BytesIO(await GanttReportService(db).generate_xlsx("F"))).active
        row = next(r for r in range(1, sheet.max_row + 1) if sheet.cell(row=r, column=1).value == "Cura do Concreto")
        striped = [c for c in range(7, sheet.max_column + 1) if sheet.cell(row=row, column=c).fill.fill_type == "darkUp"]
        assert len(striped) == 2                                        # os 2 dias da pausa atual
        legend = [sheet.cell(row=r, column=1).value for r in range(1, sheet.max_row + 1)]
        assert "Pausa" in legend

    async def test_dashboard_gantt_lists_the_pauses(self, repo, db):
        from src.services.dashboard_service import DashboardService
        await self._paused(repo)
        tasks = {t.name: t for t in await DashboardService(db).get_gantt_task_with_subtasks("Base do Gerador")}
        pauses = tasks["Base do Gerador"].pauses
        assert [(p["start"], p["reason"], p["archived"]) for p in pauses] == [
            ((TODAY - timedelta(days=2)).isoformat(), "licença", True),
            ((TODAY + timedelta(days=1)).isoformat(), "chuva", False)]
